import os
import re
import sys
import time
import random
from datetime import datetime
from pathlib import Path
from dotenv import load_dotenv
import openpyxl
from playwright.sync_api import sync_playwright

# Carrega variáveis de ambiente do ficheiro .env
load_dotenv()

BEECROWD_EMAIL = os.getenv("BEECROWD_EMAIL")
BEECROWD_PASSWORD = os.getenv("BEECROWD_PASSWORD")
BEECROWD_URL = os.getenv("BEECROWD_URL", "https://judge.beecrowd.com/pt/login")
SOLUTIONS_DIR = os.getenv("SOLUTIONS_DIR", "./")
EXCEL_FILE = os.getenv("EXCEL_FILE", "beecrowd_problems_shared.xlsx")
USER_DATA_DIR = os.getenv("USER_DATA_DIR", "./browser_session")


def apply_stealth(context) -> None:
    """Injeta scripts de ocultação para mitigar deteções básicas de automação."""
    stealth_js = """
    Object.defineProperty(navigator, 'webdriver', { get: () => undefined });
    window.navigator.chrome = { runtime: {} };
    Object.defineProperty(navigator, 'languages', { get: () => ['pt-BR', 'pt', 'en-US', 'en'] });
    Object.defineProperty(navigator, 'plugins', { get: () => [1, 2, 3, 4, 5] });
    """
    context.add_init_script(stealth_js)


def validate_environment() -> None:
    """Valida a existência das credenciais e ficheiros essenciais."""
    if not BEECROWD_EMAIL or not BEECROWD_PASSWORD:
        print("Erro: BEECROWD_EMAIL e BEECROWD_PASSWORD devem estar definidos no ficheiro .env.")
        sys.exit(1)

    if not os.path.exists(EXCEL_FILE):
        print(f"Erro: O ficheiro Excel '{EXCEL_FILE}' não foi encontrado no diretório atual.")
        sys.exit(1)


def is_captcha_blocking(page) -> bool:
    """Verifica se há um CAPTCHA visível e interativo bloqueando a execução na tela."""
    try:
        return page.evaluate("""
            () => {
                const selectors = [
                    '#cf-turnstile',
                    '.g-recaptcha',
                    '#hcaptcha',
                    'iframe[src*="turnstile"]',
                    'iframe[src*="captcha"]',
                    'iframe[src*="challenges.cloudflare.com"]'
                ];
                for (let sel of selectors) {
                    const els = document.querySelectorAll(sel);
                    for (let el of els) {
                        const rect = el.getBoundingClientRect();
                        const style = window.getComputedStyle(el);
                        if (
                            rect.width > 50 && 
                            rect.height > 50 && 
                            style.visibility !== 'hidden' && 
                            style.display !== 'none' && 
                            style.opacity !== '0'
                        ) {
                            return true;
                        }
                    }
                }
                return false;
            }
        """)
    except Exception:
        return False


def load_excel_problem_map(excel_path: str):
    """Mapeia os problemas e colunas da planilha Excel."""
    wb = openpyxl.load_workbook(excel_path)
    sheet = wb.active

    headers = [cell.value for cell in sheet[1]]

    id_col = headers.index("ID") + 1 if "ID" in headers else 1
    category_col = headers.index("Category") + 1 if "Category" in headers else None
    lang_col = headers.index("Language") + 1 if "Language" in headers else None
    approved_col = headers.index("BeeCrowd Approved") + 1 if "BeeCrowd Approved" in headers else None
    just_col = headers.index("Justification") + 1 if "Justification" in headers else None

    start_time_col = headers.index("Start Time") + 1 if "Start Time" in headers else None
    end_time_col = headers.index("End Time") + 1 if "End Time" in headers else None
    duration_col = headers.index("Execution Time") + 1 if "Execution Time" in headers else None

    problem_map = {}
    for row in range(2, sheet.max_row + 1):
        prob_id = sheet.cell(row=row, column=id_col).value
        if prob_id is not None:
            try:
                prob_id_int = int(prob_id)
                category = sheet.cell(row=row, column=category_col).value if category_col else None
                approved = sheet.cell(row=row, column=approved_col).value if approved_col else None
                problem_map[prob_id_int] = {
                    "row": row,
                    "category": str(category) if category else "",
                    "approved": str(approved) if approved else "",
                }
            except ValueError:
                continue

    wb.close()
    return problem_map, {
        "id_col": id_col,
        "category_col": category_col,
        "lang_col": lang_col,
        "approved_col": approved_col,
        "just_col": just_col,
        "start_time_col": start_time_col,
        "end_time_col": end_time_col,
        "duration_col": duration_col,
    }


def update_excel_row(
    excel_path: str,
    cols_config: dict,
    problem_id: int,
    row_num: int,
    language_used: str,
    is_approved: bool,
    verdict_text: str,
    start_time: datetime = None,
    end_time: datetime = None,
    duration_seconds: float = None,
) -> None:
    """Atualiza o resultado, horários de início/fim e tempo na coluna Execution Time."""
    wb = openpyxl.load_workbook(excel_path)
    sheet = wb.active

    if cols_config["lang_col"]:
        sheet.cell(row=row_num, column=cols_config["lang_col"]).value = language_used

    if cols_config["approved_col"]:
        sheet.cell(row=row_num, column=cols_config["approved_col"]).value = "Yes" if is_approved else "No"

    if cols_config["just_col"]:
        sheet.cell(row=row_num, column=cols_config["just_col"]).value = "" if is_approved else verdict_text

    max_col = sheet.max_column
    if not cols_config["start_time_col"]:
        max_col += 1
        cols_config["start_time_col"] = max_col
        sheet.cell(row=1, column=cols_config["start_time_col"]).value = "Start Time"

    if not cols_config["end_time_col"]:
        max_col += 1
        cols_config["end_time_col"] = max_col
        sheet.cell(row=1, column=cols_config["end_time_col"]).value = "End Time"

    if not cols_config["duration_col"]:
        max_col += 1
        cols_config["duration_col"] = max_col
        sheet.cell(row=1, column=cols_config["duration_col"]).value = "Execution Time"

    if start_time:
        sheet.cell(row=row_num, column=cols_config["start_time_col"]).value = start_time.strftime("%Y-%m-%d %H:%M:%S")

    if end_time:
        sheet.cell(row=row_num, column=cols_config["end_time_col"]).value = end_time.strftime("%Y-%m-%d %H:%M:%S")

    if duration_seconds is not None:
        sheet.cell(row=row_num, column=cols_config["duration_col"]).value = round(duration_seconds, 2)

    wb.save(excel_path)
    wb.close()

    time_info = f" | Início: {start_time.strftime('%H:%M:%S') if start_time else 'N/A'} | Fim: {end_time.strftime('%H:%M:%S') if end_time else 'N/A'} | Execution Time: {round(duration_seconds, 2) if duration_seconds is not None else 'N/A'}s"
    print(f"[Problema {problem_id}] Excel atualizado -> Aprovado: {'Yes' if is_approved else 'No'} | Justificação: '{verdict_text if not is_approved else ''}'{time_info}")


def get_latest_run_id(page) -> int:
    """Mede o ID numérico do topo da página /runs para controlo de submissões."""
    try:
        runs_url = "https://judge.beecrowd.com/pt/runs"
        if "/runs" not in page.url or "/runs/code/" in page.url:
            page.goto(runs_url, wait_until="commit", timeout=12000)

        page.wait_for_selector("table tbody tr", timeout=4000)

        run_id_str = page.evaluate(
            """
            () => {
                const firstRow = document.querySelector("table tbody tr");
                if (!firstRow) return null;
                const firstCell = firstRow.querySelector("td");
                return firstCell ? firstCell.innerText.trim() : null;
            }
        """
        )
        if run_id_str and run_id_str.isdigit():
            return int(run_id_str)
    except Exception:
        pass
    return 0


def login_to_beecrowd(page) -> bool:
    """Navega e efetua o login no Beecrowd se a sessão não estiver ativa."""
    print(f"A navegar para {BEECROWD_URL}...")

    try:
        page.goto(BEECROWD_URL, wait_until="commit", timeout=30000)
    except Exception as err:
        print(f"Aviso de carregamento inicial: {err}. A tentar novamente...")
        time.sleep(2)
        try:
            page.goto(BEECROWD_URL, wait_until="commit", timeout=30000)
        except Exception as retry_err:
            print(f"Erro ao aceder a {BEECROWD_URL}: {retry_err}")
            return False

    time.sleep(2)

    if "/login" not in page.url:
        print("Sessão já iniciada no navegador!")
        return True

    try:
        print("A aguardar formulário de login...")
        page.wait_for_selector('input[name="email"]', timeout=15000)

        print("A inserir credenciais...")
        page.fill('input[name="email"]', BEECROWD_EMAIL)
        page.fill('input[name="password"]', BEECROWD_PASSWORD)

        print("A submeter credenciais...")
        page.click('button[type="submit"], input[type="submit"]')

        time.sleep(3)

        if "/login" not in page.url:
            print("Login efetuado com sucesso!")
            return True
        else:
            print("A aguardar conclusão do login...")
            page.wait_for_url(lambda url: "/login" not in url, timeout=0)
            print("Login efetuado com sucesso!")
            return True

    except Exception as error:
        print(f"Notificação do processo de login: {error}")
        return "/login" not in page.url


def select_language(page, file_ext: str, category: str) -> str:
    """Seleciona a linguagem adequada (Java ou PostgreSQL) no dropdown."""
    is_sql = (file_ext.lower() == ".sql") or (category.upper() == "SQL")

    selected_name = page.evaluate(
        """
        (isSql) => {
            const select = document.querySelector("select[name='language_id'], select#language-id");
            if (!select) return null;

            let targetVal = null;
            let targetText = null;

            for (let opt of select.options) {
                const text = opt.text.toLowerCase();
                if (isSql) {
                    if (text.includes("postgresql") || text.includes("postgres") || text.includes("sql")) {
                        targetVal = opt.value;
                        targetText = opt.text;
                        break;
                    }
                } else {
                    if (text.includes("java")) {
                        targetVal = opt.value;
                        targetText = opt.text;
                        break;
                    }
                }
            }

            if (targetVal) {
                select.value = targetVal;
                select.dispatchEvent(new Event('change', { bubbles: true }));

                if (select.selectize) {
                    select.selectize.setValue(targetVal);
                }
                return targetText;
            }

            return null;
        }
    """,
        is_sql,
    )

    if selected_name:
        print(f"Linguagem selecionada: {selected_name}")
        return selected_name
    else:
        fallback = "PostgreSQL" if is_sql else "Java"
        print(f"Linguagem não encontrada diretamente; a assumir o valor por defeito ({fallback}).")
        return fallback


def check_is_accepted(verdict_str: str) -> bool:
    """Avalia estritamente o veredito da submissão."""
    s = verdict_str.strip().lower()

    rejection_terms = [
        "errada", "wrong", "erro", "error", "excedido", 
        "exceeded", "presentation", "apresentação", "compilação", "compilation", "pending", "unknown", "queue"
    ]
    if any(term in s for term in rejection_terms):
        return False

    return "aceito" in s or "accepted" in s


def verify_new_submission(page, problem_id: str, last_known_run_id: int, max_queue_wait_seconds: int = 120) -> tuple[str, bool]:
    """
    Navega para /runs e aguarda na página enquanto o status for 'queue' ou 'processing'
    até sair o resultado final do julgamento. Se surgir CAPTCHA, aborta e avisa.
    """
    runs_url = "https://judge.beecrowd.com/pt/runs"
    print(f"[Problema {problem_id}] A verificar registo em /runs (Run ID esperado > #{last_known_run_id})...")

    start_wait = time.time()

    while time.time() - start_wait < max_queue_wait_seconds:
        try:
            if "/runs" not in page.url or "/runs/code/" in page.url:
                page.goto(runs_url, wait_until="commit", timeout=12000)
            else:
                page.reload(wait_until="commit", timeout=12000)

            if is_captcha_blocking(page):
                print(f"\n[!] [Problema {problem_id}] CAPTCHA detectado em /runs! Pulo acionado.")
                return "CAPTCHA Required - Skipped", False

            page.wait_for_selector("table tbody tr", timeout=5000)

            top_row_data = page.evaluate(
                """
                () => {
                    const firstRow = document.querySelector("table tbody tr");
                    if (!firstRow) return null;

                    const cells = Array.from(firstRow.querySelectorAll("td")).map(td => td.innerText.trim());
                    return {
                        id: cells[0] || '0',
                        problem: cells[2] || '',
                        status: cells[4] || cells[3] || 'Unknown'
                    };
                }
            """
            )

            if top_row_data and top_row_data["id"].isdigit():
                current_run_id = int(top_row_data["id"])

                if current_run_id > last_known_run_id and str(problem_id) in top_row_data["problem"]:
                    status = top_row_data["status"]
                    status_lower = status.lower()

                    if "queue" in status_lower or "processing" in status_lower or "em fila" in status_lower:
                        print(f"\r[Problema {problem_id}] Submissão em fila (Run #{current_run_id} | Status: '{status}'). Aguardando julgamento final...", end="", flush=True)
                        time.sleep(4)
                        continue

                    print(f"\n[Problema {problem_id}] Julgamento concluído! Run #{current_run_id} | Status Final: '{status}'")
                    return status, True

        except Exception as err:
            print(f"\n[Problema {problem_id}] Aviso durante verificação de envio: {err}")

        time.sleep(3)

    print(f"\n[Problema {problem_id}] Tempo limite excedido aguardando julgamento em /runs.")
    return "Not Submitted", False


def inject_code_into_editor(page, code: str) -> bool:
    """Injeta o código-fonte no Ace Editor e sincroniza o campo de formulário oculto."""
    try:
        page.wait_for_selector(".ace_editor, textarea[name='source_code'], #source-code", state="attached", timeout=10000)
        time.sleep(1)

        injected = page.evaluate(
            """
            (codeText) => {
                function triggerEvents(element) {
                    ['focus', 'keydown', 'keypress', 'input', 'keyup', 'change', 'blur'].forEach(eventType => {
                        const event = new Event(eventType, { bubbles: true, cancelable: true });
                        element.dispatchEvent(event);
                    });
                }

                const aceContainer = document.querySelector('.ace_editor');
                if (aceContainer && window.ace) {
                    const editor = ace.edit(aceContainer);
                    editor.focus();
                    editor.setValue(codeText, -1);
                    editor.getSession().setValue(codeText);

                    if (editor.renderer) editor.renderer.updateFull();

                    const aceTextarea = aceContainer.querySelector('textarea');
                    if (aceTextarea) triggerEvents(aceTextarea);

                    const sourceCodeElement = document.querySelector('#source-code') || document.querySelector('textarea[name="source_code"]');
                    if (sourceCodeElement) {
                        sourceCodeElement.value = codeText;
                        triggerEvents(sourceCodeElement);
                    }
                    return editor.getValue().length > 0;
                }

                const textarea = document.querySelector('#source-code') || document.querySelector('textarea[name="source_code"]');
                if (textarea) {
                    textarea.focus();
                    textarea.value = codeText;
                    triggerEvents(textarea);
                    return textarea.value.length > 0;
                }

                return false;
            }
        """,
            code,
        )

        if injected:
            return True

        page.click(".ace_editor, .ace_content", timeout=5000)
        time.sleep(0.5)

        page.keyboard.press("Control+End")
        page.keyboard.type(code, delay=2)
        time.sleep(1)

        content_len = page.evaluate(
            """
            () => {
                const aceContainer = document.querySelector('.ace_editor');
                if (aceContainer && window.ace) return ace.edit(aceContainer).getValue().length;
                const hiddenTextarea = document.querySelector('#source-code') || document.querySelector('textarea[name="source_code"]');
                return hiddenTextarea ? hiddenTextarea.value.length : 0;
            }
        """
        )

        return content_len > 0

    except Exception as err:
        print(f"Erro ao injetar código no editor: {err}")
        return False


def submit_solution(page, problem_id: str, code: str, file_ext: str, category: str) -> tuple[str, bool, str, datetime, datetime, float]:
    """
    Processa o envio da solução e aguarda o veredito no /runs.
    Retorna (status, is_accepted, language_used, start_time, end_time, duration_seconds).
    Pula imediatamente se o CAPTCHA aparecer na tela.
    """
    if not code or not code.strip():
        print(f"[Problema {problem_id}] Erro: O ficheiro com o código está vazio!")
        return "Empty Code File", False, "Unknown", None, None, 0.0

    last_known_run_id = get_latest_run_id(page)
    problem_url = f"https://judge.beecrowd.com/pt/problems/view/{problem_id}"

    print(f"[Problema {problem_id}] A navegar para {problem_url}...")
    try:
        page.goto(problem_url, wait_until="commit", timeout=20000)
    except Exception as e:
        print(f"[Problema {problem_id}] Aviso na navegação: {e}")

    time.sleep(2)

    try:
        language_used = select_language(page, file_ext, category)

        print(f"[Problema {problem_id}] A injetar o código ({len(code)} bytes)...")
        inserted = inject_code_into_editor(page, code)

        if not inserted:
            print(f"[Problema {problem_id}] O código não foi preenchido corretamente no formulário. A abortar envio.")
            return "Empty Form Error", False, language_used, None, None, 0.0

        print(f"[Problema {problem_id}] Código verificado no formulário com sucesso!")
        time.sleep(1)

        print(f"[Problema {problem_id}] A clicar no botão de submissão...")
        start_time = datetime.now()

        submit_btn = page.query_selector('button[type="submit"], input[type="submit"], #btn-submit')
        if submit_btn:
            submit_btn.click()
        else:
            page.evaluate("const btn = document.querySelector('button[type=\"submit\"], input[type=\"submit\"]'); if (btn) btn.click();")

        time.sleep(3)

        if is_captcha_blocking(page):
            print(f"[!] [Problema {problem_id}] CAPTCHA exigido após submissão! Pulando para o próximo problema...")
            return "CAPTCHA Required - Skipped", False, language_used, start_time, datetime.now(), 0.0

        status, is_submitted = verify_new_submission(page, problem_id, last_known_run_id)
        end_time = datetime.now()
        duration_seconds = (end_time - start_time).total_seconds()

        if is_submitted:
            is_accepted = check_is_accepted(status)
            return status, is_accepted, language_used, start_time, end_time, duration_seconds
        else:
            return status if status != "Not Submitted" else "Submission Retained / Unconfirmed", False, language_used, start_time, end_time, 0.0

    except Exception as err:
        print(f"[Problema {problem_id}] Erro durante a submissão: {err}")
        return f"Error: {err}", False, "Unknown", None, None, 0.0


def find_solutions(base_dir: str, problem_map: dict):
    """Localiza e ordena numericamente os ficheiros de solução correspondentes aos IDs mapeados no Excel."""
    solutions = []
    base_path = Path(base_dir)

    valid_extensions = {".java", ".sql"}

    for root, _, files in os.walk(base_path):
        for file in files:
            ext = Path(file).suffix.lower()
            if ext not in valid_extensions:
                continue

            full_path = Path(root) / file

            folder_match = re.search(r"(?:problem_)?(\d+)", Path(root).name, re.IGNORECASE)
            file_match = re.search(r"(\d+)", file)

            problem_id = None
            if folder_match:
                problem_id = folder_match.group(1)
            elif file_match:
                problem_id = file_match.group(1)

            if problem_id:
                prob_int = int(problem_id)
                excel_info = problem_map.get(prob_int, {})

                solutions.append(
                    {
                        "problem_id": problem_id,
                        "prob_int": prob_int,
                        "filepath": full_path,
                        "extension": ext,
                        "category": excel_info.get("category", ""),
                        "row": excel_info.get("row"),
                        "approved": excel_info.get("approved", ""),
                    }
                )

    solutions.sort(key=lambda x: x["prob_int"])
    return solutions


def process_solution_files(page, problem_map: dict, cols_config: dict) -> None:
    """Itera sobre todas as soluções pendentes e atualiza a planilha."""
    solutions = find_solutions(SOLUTIONS_DIR, problem_map)

    if not solutions:
        print(f"Nenhum ficheiro de solução encontrado em '{SOLUTIONS_DIR}'.")
        return

    print(f"Foram encontrados {len(solutions)} ficheiro(s) de solução para processar.")

    for item in solutions:
        if item["approved"].strip().lower() == "yes":
            print(f"\n[Problema {item['problem_id']}] Já se encontra marcado como 'Yes' no Excel. A ignorar.")
            continue

        with open(item["filepath"], "r", encoding="utf-8") as f:
            code = f.read()

        status, is_accepted, language_used, start_time, end_time, duration_seconds = submit_solution(
            page=page,
            problem_id=item["problem_id"],
            code=code,
            file_ext=item["extension"],
            category=item["category"],
        )

        if item["row"]:
            update_excel_row(
                excel_path=EXCEL_FILE,
                cols_config=cols_config,
                problem_id=item["prob_int"],
                row_num=item["row"],
                language_used=language_used,
                is_approved=is_accepted,
                verdict_text=status,
                start_time=start_time,
                end_time=end_time,
                duration_seconds=duration_seconds,
            )

        if status == "CAPTCHA Required - Skipped":
            pause_time = 5.0
            print(f"[Problema {item['problem_id']}] Pulado por CAPTCHA. Pausa de {pause_time}s antes do próximo...")
        else:
            pause_time = random.uniform(35, 42)
            print(f"A aguardar {round(pause_time, 1)}s (cooldown do Beecrowd) antes da próxima submissão...")

        time.sleep(pause_time)


def main() -> None:
    """Ponto de entrada do script."""
    validate_environment()

    problem_map, cols_config = load_excel_problem_map(EXCEL_FILE)
    print(f"Foram carregadas {len(problem_map)} entradas do ficheiro {EXCEL_FILE}.")

    with sync_playwright() as p:
        print("A iniciar o Firefox com contexto persistente...")
        context = p.firefox.launch_persistent_context(
            user_data_dir=USER_DATA_DIR,
            headless=False
        )

        apply_stealth(context)

        page = context.pages[0] if context.pages else context.new_page()

        if login_to_beecrowd(page):
            process_solution_files(page, problem_map, cols_config)

        print("\nProcesso concluído. A fechar o navegador...")
        time.sleep(3)
        context.close()


if __name__ == "__main__":
    main()
